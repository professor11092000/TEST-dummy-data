"""
Migrate old_tracker.xlsx to new_template.xlsx locally.
Scan Relevancy supports numeric values (1, 2, 3, 4...) as well as text labels.
The scan number is used to group rows from the same application scan round.

Run:
    pip install openpyxl
    python migrate_tracker.py

Keep both input files in the same folder as this script.
The script writes migrated_tracker.xlsx and never uploads data anywhere.
"""

from pathlib import Path
from datetime import date, datetime
import re
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

OLD_FILE = Path("old_tracker.xlsx")
TEMPLATE_FILE = Path("new_template.xlsx")
OUTPUT_FILE = Path("migrated_tracker.xlsx")

# Confirm these values with your approved SLA policy.
SLA_DAYS = {"critical": None, "high": 45, "medium": 90, "low": 180}
TODAY = date.today()


def norm(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def headers(ws):
    return {norm(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1)
            if ws.cell(1, c).value is not None}


def col(hmap, *names):
    for name in names:
        if norm(name) in hmap:
            return hmap[norm(name)]
    return None


def val(row, hmap, *names, default=None):
    c = col(hmap, *names)
    return row[c - 1] if c and c - 1 < len(row) else default


def put(ws, hmap, rownum, header, value):
    c = col(hmap, header)
    if c:
        ws.cell(rownum, c, value)


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not v:
        return None
    for fmt in ("%d-%b-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            pass
    return None


def count_instances(v):
    try:
        return max(1, int(float(v)))
    except (TypeError, ValueError):
        return 1


def sev_and_title(severity_value, title_value):
    """Defensively handle rows where a finding name was entered in Severity."""
    allowed = {"critical", "high", "medium", "low", "informational", "info"}
    sev_text = str(severity_value or "").strip()
    title_text = str(title_value or "").strip()
    if sev_text.lower() not in allowed and not title_text:
        return "", sev_text
    return sev_text, title_text


def scan_type(assessment):
    a = str(assessment or "").lower()
    if "dast" in a or "dynamic" in a:
        return "Dynamic"
    if any(x in a for x in ("sast", "sca", "iac", "container", "static")):
        return "Static"
    return "Unknown" if a else ""


def status_map(v):
    s = str(v or "").strip().lower()
    if s in {"fixed", "closed", "resolved", "remediated"}:
        return "Closed"
    if s in {"active", "open", "reopened", "in progress"}:
        return "Open"
    return str(v or "").strip()


def main():
    if not OLD_FILE.exists() or not TEMPLATE_FILE.exists():
        raise FileNotFoundError(
            "Place old_tracker.xlsx and new_template.xlsx beside this script."
        )

    old_wb = load_workbook(OLD_FILE, data_only=True)
    if "Application Inventory" not in old_wb.sheetnames or "Vulnerability Tracker" not in old_wb.sheetnames:
        raise ValueError("Old workbook must contain 'Application Inventory' and 'Vulnerability Tracker'.")

    inv_ws = old_wb["Application Inventory"]
    findings_ws = old_wb["Vulnerability Tracker"]
    ih, fh = headers(inv_ws), headers(findings_ws)

    out_wb = load_workbook(TEMPLATE_FILE)
    for sheet in ("App Inventory", "Scans", "Vulnerabilities"):
        if sheet not in out_wb.sheetnames:
            raise ValueError(f"New template is missing worksheet: {sheet}")
    app_ws, scans_ws, vuln_ws = (
        out_wb["App Inventory"], out_wb["Scans"], out_wb["Vulnerabilities"]
    )
    ah, sh, vh = headers(app_ws), headers(scans_ws), headers(vuln_ws)

    # Clear existing sample/data rows while preserving row 1 and its formatting.
    for ws in (app_ws, scans_ws, vuln_ws):
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)

    # Application Inventory: one output row per unique application key.
    apps = {}
    for row in inv_ws.iter_rows(min_row=2, values_only=True):
        if not any(x is not None for x in row):
            continue
        app_name = str(val(row, ih, "Application Name", "App Name", default="") or "").strip()
        unique_name = str(val(row, ih, "Unique Application Name", "Unique App Name", default=app_name) or app_name).strip()
        if not app_name and not unique_name:
            continue
        key = (unique_name or app_name).casefold()
        if key not in apps:
            apps[key] = {
                "id": f"APP-{len(apps)+1:04d}",
                "name": app_name or unique_name,
                "unique": unique_name or app_name,
                "source": row,
                "last_scan_date": val(row, ih, "Last Scan Date"),
                "last_report_shared_date": val(row, ih, "Report Shared Date"),
            }

    # Add any app appearing in Vulnerability Tracker but absent from inventory.
    for row in findings_ws.iter_rows(min_row=2, values_only=True):
        if not any(x is not None for x in row):
            continue
        app_name = str(val(row, fh, "Application Name", "App Name", default="") or "").strip()
        unique_name = str(val(row, fh, "Unique Application Name", "Unique App Name", default=app_name) or app_name).strip()
        if not app_name and not unique_name:
            continue
        key = (unique_name or app_name).casefold()
        if key not in apps:
            apps[key] = {"id": f"APP-{len(apps)+1:04d}", "name": app_name or unique_name,
                         "unique": unique_name or app_name, "source": None,
                         "last_scan_date": None, "last_report_shared_date": None}

    # Write app inventory fields only where matching template headers exist.
    for i, app in enumerate(apps.values(), start=2):
        put(app_ws, ah, i, "App ID", app["id"])
        put(app_ws, ah, i, "App Name", app["name"])
        put(app_ws, ah, i, "Unique App Name", app["unique"])
        source = app["source"]
        if source:
            put(app_ws, ah, i, "BU", val(source, ih, "Application Group"))
            put(app_ws, ah, i, "App Owner", val(source, ih, "Primary Application Owner"))
            put(app_ws, ah, i, "Criticality", val(source, ih, "App Risk Category"))
            put(app_ws, ah, i, "App Type", val(source, ih, "Application Type"))
            put(app_ws, ah, i, "Repo URL", val(source, ih, "Repo URL"))
            put(app_ws, ah, i, "Test Env URL", val(source, ih, "Test Env URL"))

    # A scan is one application + scan relevancy label. Do NOT split by Assessment Type:
    # one scan can contain DAST/SAST/SCA/IaC findings. Create only scan groups evidenced
    # by at least one row in the source; do not invent scans for missing periods.
    finding_rows = []
    scan_groups = {}
    for row in findings_ws.iter_rows(min_row=2, values_only=True):
        if not any(x is not None for x in row):
            continue
        app_name = str(val(row, fh, "Application Name", "App Name", default="") or "").strip()
        unique_name = str(val(row, fh, "Unique Application Name", "Unique App Name", default=app_name) or app_name).strip()
        app_key = (unique_name or app_name).casefold()
        if app_key not in apps:
            continue
        relevancy = str(val(row, fh, "Scan Relevancy", default="") or "").strip()
        scan_date_raw = val(row, fh, "Scan Date")
        # Scan date is blank in some source rows. Do not substitute the application's
        # current/last scan date for historical scans.
        scan_date = as_date(scan_date_raw)
        # Numeric scan relevancy identifies the scan round for an application.
        # Group by application + scan number, not Assessment Type or date, so
        # DAST/SAST/SCA findings in the same round share one Scan ID.
        group_key = (app_key, relevancy.casefold())
        if group_key not in scan_groups:
            scan_groups[group_key] = {
                "id": f"SCAN-{len(scan_groups)+1:06d}", "app_key": app_key,
                "relevancy": relevancy, "date": scan_date, "date_raw": scan_date_raw,
                "is_latest": str(val(row, fh, "Latest Scan?", default="") or "").strip().lower() == "latest scan",
                "report_shared_date": as_date(val(row, fh, "Report Shared Date")),
                "rows": [],
            }
        scan_groups[group_key]["rows"].append(row)
        finding_rows.append((row, group_key))

    # Write one row per scan group, regardless of how many Assessment Types are in it.
    for i, group in enumerate(scan_groups.values(), start=2):
        app = apps[group["app_key"]]
        expanded = []
        for row in group["rows"]:
            raw_sev = val(row, fh, "Vulnerability Severity", "Severity")
            raw_title = val(row, fh, "Vulnerability Name", "Finding Title")
            severity, title = sev_and_title(raw_sev, raw_title)
            expanded.extend([severity] * count_instances(val(row, fh, "Instance Count", "Instance")))
        put(scans_ws, sh, i, "SCAN ID", group["id"])
        put(scans_ws, sh, i, "APP ID", app["id"])
        put(scans_ws, sh, i, "APP NAME", app["name"])
        inferred_latest_scan_date = app.get("last_scan_date") if group["is_latest"] else None
        scan_date_value = group["date"] or group["date_raw"] or inferred_latest_scan_date
        report_date_value = group["report_shared_date"] or (
            app.get("last_report_shared_date") if group["is_latest"] else None
        )
        put(scans_ws, sh, i, "DATE", scan_date_value)
        put(scans_ws, sh, i, "Scan Date", scan_date_value)
        put(scans_ws, sh, i, "Report Shared Date", report_date_value)
        put(scans_ws, sh, i, "Scan Completion Date", scan_date_value)
        # Assessment type is deliberately not used to split scan groups.
        put(scans_ws, sh, i, "Scan Type", "Mixed" if len({
            str(val(r, fh, "Assessment Type", default="") or "").strip().upper()
            for r in group["rows"]
        }) > 1 else scan_type(val(group["rows"][0], fh, "Assessment Type")))
        # Scan Relevancy can be numeric (1, 2, 3, 4...) or text
        # ("First Scan", "Second Scan", etc.). Number 1 means first/new scan;
        # numbers greater than 1 mean a rescan.
        rel = str(group["relevancy"] or "").strip().lower()
        numeric_match = re.fullmatch(r"(?:scan\s*)?(\d+)", rel)
        if numeric_match:
            scan_number = int(numeric_match.group(1))
            exec_type = "New" if scan_number == 1 else "Rescan"
        elif "first" in rel:
            exec_type = "New"
        elif rel:
            exec_type = "Rescan"
        else:
            exec_type = "Unknown"
        put(scans_ws, sh, i, "Execution Type", exec_type)
        put(scans_ws, sh, i, "Status", "Completed")
        put(scans_ws, sh, i, "Total Findings", len(expanded))
        for bucket in ("Critical", "High", "Medium", "Low"):
            put(scans_ws, sh, i, bucket,
                sum(1 for s in expanded if str(s or "").strip().lower() == bucket.lower()))

    scan_id_by_key = {key: group["id"] for key, group in scan_groups.items()}

    # Determine New vs Recurring by comparing each scan round with the previous
    # available scan round for the same application. Match on normalized finding
    # title + Assessment Type (e.g., DAST/SAST/SCA). The first scan is all New.
    def finding_identity(source_row):
        raw_sev = val(source_row, fh, "Vulnerability Severity", "Severity")
        raw_title = val(source_row, fh, "Vulnerability Name", "Finding Title")
        _severity, title = sev_and_title(raw_sev, raw_title)
        assessment = str(val(source_row, fh, "Assessment Type", "Finding From", default="") or "").strip()
        return (re.sub(r"[^a-z0-9]+", "", str(title or "").lower()),
                re.sub(r"[^a-z0-9]+", "", assessment.lower()))

    def scan_number(label):
        match = re.fullmatch(r"\s*(?:scan\s*)?(\d+)\s*", str(label or ""), flags=re.IGNORECASE)
        if match:
            return int(match.group(1))
        # Retain compatibility with old textual labels if any remain.
        words = {"first": 1, "second": 2, "third": 3, "fourth": 4, "fifth": 5,
                 "sixth": 6, "seventh": 7, "eighth": 8, "ninth": 9, "tenth": 10}
        text = str(label or "").strip().lower()
        for word, number in words.items():
            if word in text:
                return number
        return None

    app_scan_order = {}
    for group_key, group in scan_groups.items():
        app_scan_order.setdefault(group["app_key"], []).append(
            (scan_number(group["relevancy"]), group_key, group)
        )
    recurring_by_row_id = {}
    for app_key, entries in app_scan_order.items():
        entries.sort(key=lambda item: (item[0] is None, item[0] if item[0] is not None else 10**9))
        previous_findings = set()
        for number, group_key, group in entries:
            current_findings = {finding_identity(source_row) for source_row in group["rows"]}
            for source_row in group["rows"]:
                identity = finding_identity(source_row)
                # If there is no earlier scan round, classify as New.
                recurring_by_row_id[id(source_row)] = (
                    "Recurring" if number is not None and number > 1 and identity in previous_findings
                    else "New"
                )
            previous_findings = current_findings

    # One output Vulnerabilities row per instance. Each row points to the same Scan ID
    # for all findings from that app/scan, including different assessment types.
    finding_id = 0
    for row, group_key in finding_rows:
        app_key = group_key[0]
        app = apps[app_key]
        group = scan_groups[group_key]
        raw_sev = val(row, fh, "Vulnerability Severity", "Severity")
        raw_title = val(row, fh, "Vulnerability Name", "Finding Title")
        severity, title = sev_and_title(raw_sev, raw_title)
        discovery_raw = val(row, fh, "First Identified Date", "Discovery Date")
        discovery = as_date(discovery_raw)
        closed_raw = val(row, fh, "Closed date", "Closed Date", "Closure Date")
        closed = as_date(closed_raw)
        status = status_map(val(row, fh, "Remediation Status", "Status"))
        assessment = val(row, fh, "Assessment Type", "Finding From")
        due = discovery.fromordinal(discovery.toordinal() + SLA_DAYS[severity.lower()]) \
              if discovery and severity.lower() in SLA_DAYS and SLA_DAYS[severity.lower()] else None
        if discovery:
            end = closed if status == "Closed" and closed else TODAY
            days_open = max(0, (end - discovery).days)
        else:
            days_open = None
        if due is None:
            sla_status = "N/A"
        elif status == "Closed":
            sla_status = "Met" if closed and closed <= due else ("Breached" if closed else "N/A")
        else:
            sla_status = "Overdue" if TODAY > due else "Within SLA"

        for _ in range(count_instances(val(row, fh, "Instance Count", "Instance"))):
            finding_id += 1
            out_row = finding_id + 1
            put(vuln_ws, vh, out_row, "Finding ID", f"FIND-{finding_id:07d}")
            put(vuln_ws, vh, out_row, "App ID", app["id"])
            put(vuln_ws, vh, out_row, "Scan ID", scan_id_by_key[group_key])
            put(vuln_ws, vh, out_row, "Finding Title", title)
            put(vuln_ws, vh, out_row, "Severity", severity)
            put(vuln_ws, vh, out_row, "Vulnerability Type", recurring_by_row_id.get(id(row), "New"))
            put(vuln_ws, vh, out_row, "Category", "")
            put(vuln_ws, vh, out_row, "Discovery Date", discovery or discovery_raw)
            put(vuln_ws, vh, out_row, "Finding From", assessment)
            put(vuln_ws, vh, out_row, "Scan Type", scan_type(assessment))
            put(vuln_ws, vh, out_row, "SLA Days", SLA_DAYS.get(severity.lower()))
            put(vuln_ws, vh, out_row, "SLA Due Date", due)
            put(vuln_ws, vh, out_row, "Days Open", days_open)
            put(vuln_ws, vh, out_row, "Status", status)
            put(vuln_ws, vh, out_row, "SLA Status", sla_status)
            put(vuln_ws, vh, out_row, "Assigned To", "")
            put(vuln_ws, vh, out_row, "Closure Date", closed or closed_raw)

    # Keep template headers and formatting, apply date formats, filters, freeze panes.
    for ws in (app_ws, scans_ws, vuln_ws):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                header = str(ws.cell(1, cell.column).value or "").lower()
                if isinstance(cell.value, (date, datetime)) and "date" in header:
                    cell.number_format = "dd-mmm-yyyy"
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    out_wb.save(OUTPUT_FILE)
    print("Migration complete:", OUTPUT_FILE.resolve())
    print("Applications:", len(apps))
    print("Scan groups created from scan rows actually present:", len(scan_groups))
    print("Expanded vulnerability instance rows:", finding_id)
    print("Note: missing latest-scan dates can fall back to App Inventory's Last Scan Date.")
    print("      Historical scan dates, first identified dates, and closure dates are not invented.")
    print("Vulnerability Type is classified as New or Recurring by comparing normalized")
    print("finding title + Assessment Type with the previous available scan round for that app.")
    print("The first scan is New by default; recurrence matching is exact after normalization.") 
    print("Note: if a source row has a non-standard severity and blank vulnerability name,")
    print("      the script treats the source 'Vulnerability Severity' value as the finding title.")


if __name__ == "__main__":
    main()
