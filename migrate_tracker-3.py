#!/usr/bin/env python3
"""
Combined Static + Dynamic tracker migration.

Put these files beside this script:
    static.xlsx
    dynamic.xlsx
    new_template.xlsx

Run:
    python migrate_tracker.py

Output:
    migrated_tracker.xlsx
"""

from pathlib import Path
from datetime import datetime, date, timedelta
from copy import copy
import re
from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
STATIC_FILE = BASE / "static.xlsx"
DYNAMIC_FILE = BASE / "dynamic.xlsx"
TEMPLATE_FILE = BASE / "new_template.xlsx"
OUTPUT_FILE = BASE / "migrated_tracker.xlsx"


def txt(v):
    return "" if v is None else str(v).strip()


def norm(v):
    return re.sub(r"\s+", " ", txt(v)).strip().lower()


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, (datetime, date)):
        return v
    s = txt(v)
    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%d/%m/%Y",
                "%m/%d/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return v


def to_int(v, default=1):
    try:
        n = int(float(txt(v)))
        return n if n > 0 else default
    except Exception:
        return default


def get_headers(ws):
    return {
        norm(c.value): c.column
        for c in ws[1]
        if txt(c.value)
    }


def find_col(h, aliases):
    if isinstance(aliases, str):
        aliases = [aliases]
    for alias in aliases:
        c = h.get(norm(alias))
        if c is not None:
            return c
    return None


def get_value(row, h, aliases, default=""):
    c = find_col(h, aliases)
    if c is None:
        return default
    # row contains values, not cell objects.
    return row[c - 1]


def set_value(ws, h, row_no, aliases, value):
    c = find_col(h, aliases)
    if c is None:
        return False
    ws.cell(row_no, c).value = value
    return True


def copy_style(ws, source_row, target_row):
    if source_row < 2:
        return
    for c in range(1, ws.max_column + 1):
        src = ws.cell(source_row, c)
        dst = ws.cell(target_row, c)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)


def append_row(ws):
    return max(2, ws.max_row + 1)


# ============================================================
# OLD TRACKER COLUMN ALIASES
# ============================================================

OLD = {
    "app": ["Application Name", "App Name", "Application"],
    "unique_app": ["Unique Application Name", "Unique App Name"],
    "scan_no": ["Scan Relevancy", "Scan Number", "Scan No", "Scan"],
    "latest": ["Latest Scan?", "Latest Scan", "Is Latest Scan"],
    "first_date": ["First Identified Date", "Discovery Date", "First Identified"],
    "scan_date": ["Scan Date", "Date", "Scan Completion Date"],
    "report_date": ["Report Shared Date", "Report Date"],
    "closed_date": ["Closed date", "Closed Date", "Closure Date"],
    "severity": ["Vulnerability Severity", "Severity"],
    "finding": ["Vulnerability Name", "Finding Title", "Finding Name"],
    "status": ["Remediation Status", "Status"],
    "assessment": ["Assessment Type", "Finding From", "Scan Type", "Execution Type"],
    "comment": ["Comment (if any)", "Comment", "Remarks"],
    "instances": ["Instance Count", "Instances", "Instance"],
    "url": ["App URL", "Application URL", "URL", "URL Link"],
    "bu": ["Business Unit", "BU", "First BU"],
    "criticality": ["Application Criticality", "App Criticality", "Criticality"],
    "frequency": ["Scan Frequency", "Frequency"],
    "tool": ["Tool", "Scan Tool"],
    "execution": ["Execution Type"],
}


# ============================================================
# NEW TEMPLATE COLUMN ALIASES
# ============================================================

NEW = {
    "app_id": ["App ID"],
    "app_name": ["Application Name", "App Name"],
    "unique_app": ["Unique Application Name", "Unique App Name"],
    "app_url": ["App URL", "Application URL", "URL", "URL Link"],
    "bu": ["Business Unit", "BU", "First BU"],
    "criticality": ["Application Criticality", "App Criticality", "Criticality"],
    "frequency": ["Scan Frequency", "Frequency"],
    "scan_type": ["Scan Type"],

    "scan_id": ["Scan ID"],
    "date": ["Date", "Scan Date"],
    "tool": ["Tool", "Scan Tool"],
    "execution": ["Execution Type"],
    "scan_status": ["Status"],
    "completion_date": ["Scan Completion Date"],
    "total": ["Total Findings"],
    "critical": ["Critical"],
    "high": ["High"],
    "medium": ["Medium"],
    "low": ["Low"],
    "latest": ["Latest Scan?", "Latest Scan"],

    "vul_id": ["Vul ID", "Vulnerability ID"],
    "finding_title": ["Finding Title", "Vulnerability Name", "Finding Name"],
    "severity": ["Severity", "Vulnerability Severity"],
    "vul_type": ["Vulnerability Type", "Type"],
    "category": ["Category"],
    "discovery": ["Discovery Date", "First Identified Date"],
    "finding_from": ["Finding From"],
    "sla_days": ["SLA Days"],
    "sla_due": ["SLA Due Date"],
    "days_open": ["Days Open"],
    "status": ["Status", "Remediation Status"],
    "closed_date": ["Closed Date", "Closed date", "Closure Date"],
    "sla_status": ["SLA Status"],
    "comment": ["Comment", "Comment (if any)", "Remarks"],
    "app_url_v": ["App URL", "Application URL", "URL"],
}


SLA_DAYS = {
    "critical": 30,
    "high": 45,
    "medium": 90,
    "low": 180,
}


def normalize_latest(v):
    n = norm(v)
    if n in {"latest", "latest scan", "yes", "true", "y"}:
        return "Latest Scan"
    if n in {"older", "older scan", "no", "false", "n"}:
        return "Older Scan"
    return txt(v)


def new_status(v):
    n = norm(v)
    if n == "active" or n == "open":
        return "Open"
    if n == "fixed" or n == "closed":
        return "Closed"
    return txt(v)


def calculate_sla(severity, discovery_date, status, closed_date):
    days = SLA_DAYS.get(norm(severity))
    if not days or not discovery_date:
        return None, None, None, ""

    try:
        discovery = (
            discovery_date.date()
            if isinstance(discovery_date, datetime)
            else discovery_date
        )
        due = discovery + timedelta(days=days)

        if norm(status) == "closed":
            end_date = (
                closed_date.date()
                if isinstance(closed_date, datetime)
                else closed_date
            )
            if not end_date:
                end_date = date.today()

            age = (end_date - discovery).days
            sla_status = "Breached" if end_date > due else "Within SLA"
        else:
            age = (date.today() - discovery).days
            sla_status = "Overdue" if date.today() > due else "Within SLA"

        return days, due, age, sla_status
    except Exception:
        return days, None, None, ""


def scan_sort_key(v):
    s = txt(v)
    try:
        return (0, int(float(s)))
    except Exception:
        return (1, s.lower())


def read_source(path):
    wb = load_workbook(path, data_only=False)

    # Do NOT assume the old tracker data is on a sheet named
    # "Vulnerabilities". Some old workbooks contain a sheet with that
    # name that belongs to a different/new structure. Find the worksheet
    # whose header row actually matches the OLD tracker columns.
    required = [
        ("app", "Application Name"),
        ("scan_no", "Scan Relevancy"),
        ("latest", "Latest Scan?"),
        ("first_date", "First Identified Date"),
        ("scan_date", "Scan Date"),
        ("severity", "Vulnerability Severity"),
        ("finding", "Vulnerability Name"),
        ("status", "Remediation Status"),
        ("assessment", "Assessment Type"),
    ]

    ws = None
    h = None
    best_missing = None

    for candidate in wb.worksheets:
        candidate_h = get_headers(candidate)
        missing = [
            label for key, label in required
            if find_col(candidate_h, OLD[key]) is None
        ]

        if not missing:
            ws = candidate
            h = candidate_h
            break

        if best_missing is None or len(missing) < len(best_missing):
            best_missing = missing

    if ws is None:
        raise ValueError(
            f"Source tracker {path.name} does not contain a worksheet "
            "with the expected OLD tracker headers. Missing: "
            + ", ".join(best_missing or [label for _, label in required])
            + "\nAvailable worksheets: "
            + ", ".join(wb.sheetnames)
        )

    print(f"  Using source worksheet: {ws.title}")

    records = []

    for r in range(2, ws.max_row + 1):
        # IMPORTANT: row contains cell VALUES.
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]

        app = txt(get_value(row, h, OLD["app"]))
        if not app:
            continue

        records.append({
            "app": app,
            "unique_app": txt(get_value(row, h, OLD["unique_app"], app)) or app,
            "scan_no": txt(get_value(row, h, OLD["scan_no"])),
            "latest": normalize_latest(get_value(row, h, OLD["latest"])),
            "first_date": to_date(get_value(row, h, OLD["first_date"])),
            "scan_date": to_date(get_value(row, h, OLD["scan_date"])),
            "report_date": to_date(get_value(row, h, OLD["report_date"])),
            "closed_date": to_date(get_value(row, h, OLD["closed_date"])),
            "severity": txt(get_value(row, h, OLD["severity"])),
            "finding": txt(get_value(row, h, OLD["finding"])),
            "status": txt(get_value(row, h, OLD["status"])),
            "assessment": txt(get_value(row, h, OLD["assessment"])),
            "comment": txt(get_value(row, h, OLD["comment"])),
            "instances": to_int(get_value(row, h, OLD["instances"], 1), 1),
            "url": txt(get_value(row, h, OLD["url"])),
            "bu": txt(get_value(row, h, OLD["bu"])),
            "criticality": txt(get_value(row, h, OLD["criticality"])),
            "frequency": txt(get_value(row, h, OLD["frequency"])),
            "tool": txt(get_value(row, h, OLD["tool"])),
            "execution": txt(get_value(row, h, OLD["execution"])),
        })

    return records


def finding_from(rec, tracker_type):
    if tracker_type == "static":
        return rec["assessment"]

    a = norm(rec["assessment"])
    if a.startswith("m"):
        return "Manual Scan"
    if a.startswith("a"):
        return "Automated Scan"
    return rec["assessment"]


def migrate_one(records, wb, tracker_type, counters):
    app_ws = wb["App Inventory"]
    scan_ws = wb["Scans"]
    vul_ws = wb["Vulnerabilities"]

    ah = get_headers(app_ws)
    sh = get_headers(scan_ws)
    vh = get_headers(vul_ws)

    scan_type = "Static" if tracker_type == "static" else "Dynamic"

    # Validate using the same alias logic used when writing values.
    # This is important because the template may use "App Name" while
    # another version may use "Application Name".
    required_new = {
        "App Inventory": {
            "App ID": NEW["app_id"],
            "Application Name": NEW["app_name"],
            "Scan Type": NEW["scan_type"],
        },
        "Scans": {
            "Scan ID": NEW["scan_id"],
            "App ID": NEW["app_id"],
            "App Name": NEW["app_name"],
            "Date": NEW["date"],
            "Scan Type": NEW["scan_type"],
            "Latest Scan?": NEW["latest"],
        },
        "Vulnerabilities": {
            "Vul ID": NEW["vul_id"],
            "App ID": NEW["app_id"],
            "Scan ID": NEW["scan_id"],
            "Finding Title": NEW["finding_title"],
            "Severity": NEW["severity"],
            "Vulnerability Type": NEW["vul_type"],
            "Discovery Date": NEW["discovery"],
            "Finding From": NEW["finding_from"],
            "Scan Type": NEW["scan_type"],
            "Status": NEW["status"],
            "Latest Scan?": NEW["latest"],
        },
    }

    maps = {
        "App Inventory": ah,
        "Scans": sh,
        "Vulnerabilities": vh,
    }

    for sheet_name, needed in required_new.items():
        missing = [
            label for label, aliases in needed.items()
            if find_col(maps[sheet_name], aliases) is None
        ]
        if missing:
            raise ValueError(
                f"{sheet_name} is missing required columns: " + ", ".join(missing)
            )

    # --------------------------------------------------------
    # APP IDs
    # --------------------------------------------------------
    app_records = {}
    for rec in records:
        key = (norm(rec["app"]), tracker_type)
        if key not in app_records:
            app_records[key] = rec

    app_ids = {}

    for key, rec in app_records.items():
        app_id = f"APP-{counters['app']:06d}"
        counters["app"] += 1
        app_ids[key] = app_id

        r = append_row(app_ws)
        copy_style(app_ws, max(2, r - 1), r)

        set_value(app_ws, ah, r, NEW["app_id"], app_id)
        set_value(app_ws, ah, r, NEW["app_name"], rec["app"])
        set_value(app_ws, ah, r, NEW["unique_app"], rec["unique_app"])
        set_value(app_ws, ah, r, NEW["scan_type"], scan_type)
        set_value(app_ws, ah, r, NEW["app_url"], rec["url"])
        set_value(app_ws, ah, r, NEW["bu"], rec["bu"])
        set_value(app_ws, ah, r, NEW["criticality"], rec["criticality"])
        set_value(app_ws, ah, r, NEW["frequency"], rec["frequency"])

    # --------------------------------------------------------
    # GROUP BY APP + SCAN + TRACKER TYPE
    # --------------------------------------------------------
    scans = {}

    for rec in records:
        key = (norm(rec["app"]), rec["scan_no"], tracker_type)

        if key not in scans:
            scans[key] = {
                "app": rec["app"],
                "scan_no": rec["scan_no"],
                "latest": rec["latest"],
                "scan_date": rec["scan_date"],
                "report_date": rec["report_date"],
                "url": rec["url"],
                "tool": rec["tool"],
                "execution": rec["execution"],
                "records": [],
                "app_key": (norm(rec["app"]), tracker_type),
            }

        scans[key]["records"].append(rec)

        if rec["latest"] == "Latest Scan":
            scans[key]["latest"] = "Latest Scan"

        for field in ("scan_date", "report_date", "url", "tool", "execution"):
            if not scans[key][field] and rec[field]:
                scans[key][field] = rec[field]

    # --------------------------------------------------------
    # PREVIOUS-SCAN FINDING SET
    # --------------------------------------------------------
    by_app = {}

    for scan in scans.values():
        by_app.setdefault(scan["app_key"], []).append(scan)

    previous_findings = {}

    for app_key, scan_list in by_app.items():
        scan_list.sort(key=lambda x: scan_sort_key(x["scan_no"]))
        previous_set = set()

        for scan in scan_list:
            key = (norm(scan["app"]), scan["scan_no"], tracker_type)
            previous_findings[key] = previous_set.copy()

            current_set = {
                norm(r["finding"])
                for r in scan["records"]
                if r["finding"] and norm(r["severity"]) != "not found"
            }

            previous_set = current_set

    # --------------------------------------------------------
    # CREATE SCANS + VULNERABILITIES
    # --------------------------------------------------------
    for scan_key, scan in scans.items():

        scan_id = f"SCAN-{counters['scan']:06d}"
        counters["scan"] += 1

        app_id = app_ids[scan["app_key"]]

        # ---------------- SCANS ----------------
        scan_row = append_row(scan_ws)
        copy_style(scan_ws, max(2, scan_row - 1), scan_row)

        set_value(scan_ws, sh, scan_row, NEW["scan_id"], scan_id)
        set_value(scan_ws, sh, scan_row, ["App ID"], app_id)
        set_value(scan_ws, sh, scan_row, ["App Name", "Application Name"], scan["app"])
        set_value(scan_ws, sh, scan_row, NEW["date"], scan["scan_date"])
        set_value(scan_ws, sh, scan_row, NEW["scan_type"], scan_type)
        set_value(scan_ws, sh, scan_row, NEW["tool"], scan["tool"])
        set_value(scan_ws, sh, scan_row, NEW["execution"], scan["execution"])
        set_value(scan_ws, sh, scan_row, NEW["scan_status"], "Completed")
        set_value(scan_ws, sh, scan_row, NEW["completion_date"], scan["scan_date"])
        set_value(scan_ws, sh, scan_row, NEW["latest"], scan["latest"])

        # ---------------- VULNERABILITIES ----------------
        counts = {"critical": 0, "high": 0, "medium": 0, "low": 0, "total": 0}
        previous_set = previous_findings.get(scan_key, set())

        try:
            first_scan = int(float(scan["scan_no"])) == 1
        except Exception:
            first_scan = norm(scan["scan_no"]) in {"first scan", "first"}

        for rec in scan["records"]:

            # "Not Found" means scan occurred but there is no
            # vulnerability row.
            if norm(rec["severity"]) == "not found":
                continue

            if not rec["finding"]:
                continue

            vul_type = (
                "New"
                if first_scan
                else ("Recurring" if norm(rec["finding"]) in previous_set else "New")
            )

            source_type = finding_from(rec, tracker_type)

            # Static instance count becomes individual rows.
            # Dynamic remains one row per finding.
            copies = max(1, rec["instances"]) if tracker_type == "static" else 1

            for _ in range(copies):

                vul_id = f"FIND-{counters['finding']:06d}"
                counters["finding"] += 1

                status = new_status(rec["status"])

                sla_days, sla_due, days_open, sla_status = calculate_sla(
                    rec["severity"],
                    rec["first_date"],
                    status,
                    rec["closed_date"],
                )

                r = append_row(vul_ws)
                copy_style(vul_ws, max(2, r - 1), r)

                set_value(vul_ws, vh, r, NEW["vul_id"], vul_id)
                set_value(vul_ws, vh, r, ["App ID"], app_id)
                set_value(vul_ws, vh, r, ["Scan ID"], scan_id)
                set_value(vul_ws, vh, r, NEW["finding_title"], rec["finding"])
                set_value(vul_ws, vh, r, NEW["severity"], rec["severity"])
                set_value(vul_ws, vh, r, NEW["vul_type"], vul_type)

                # First Identified Date -> Discovery Date.
                # All repeated instances retain the same original date.
                set_value(vul_ws, vh, r, NEW["discovery"], rec["first_date"])

                set_value(vul_ws, vh, r, NEW["finding_from"], source_type)
                set_value(vul_ws, vh, r, NEW["scan_type"], scan_type)

                set_value(vul_ws, vh, r, NEW["sla_days"], sla_days)
                set_value(vul_ws, vh, r, NEW["sla_due"], sla_due)
                set_value(vul_ws, vh, r, NEW["days_open"], days_open)
                set_value(vul_ws, vh, r, NEW["status"], status)
                set_value(vul_ws, vh, r, NEW["closed_date"], rec["closed_date"])
                set_value(vul_ws, vh, r, NEW["sla_status"], sla_status)
                set_value(vul_ws, vh, r, NEW["comment"], rec["comment"])
                set_value(vul_ws, vh, r, NEW["app_url_v"], rec["url"])

                # Critical for dashboard filtering:
                # carry Latest Scan? to every vulnerability row.
                set_value(vul_ws, vh, r, NEW["latest"], scan["latest"])

                counts["total"] += 1
                sev = norm(rec["severity"])
                if sev in counts:
                    counts[sev] += 1

        # ---------------- SCAN SUMMARY ----------------
        set_value(scan_ws, sh, scan_row, NEW["total"], counts["total"])
        set_value(scan_ws, sh, scan_row, NEW["critical"], counts["critical"])
        set_value(scan_ws, sh, scan_row, NEW["high"], counts["high"])
        set_value(scan_ws, sh, scan_row, NEW["medium"], counts["medium"])
        set_value(scan_ws, sh, scan_row, NEW["low"], counts["low"])

    return {"apps": len(app_records), "scans": len(scans)}


def main():
    print("=" * 60)
    print("Starting combined migration...")
    print("=" * 60)

    for f in (STATIC_FILE, DYNAMIC_FILE, TEMPLATE_FILE):
        if not f.exists():
            raise FileNotFoundError(
                f"\nMissing file: {f.name}\n"
                f"Place {f.name} in the same folder as this script:\n"
                f"{BASE}"
            )

    print(f"Static source : {STATIC_FILE.name}")
    print(f"Dynamic source: {DYNAMIC_FILE.name}")
    print(f"Template      : {TEMPLATE_FILE.name}")
    print()

    print("Reading Static tracker...")
    static_records = read_source(STATIC_FILE)
    print(f"  Static source rows: {len(static_records)}")

    print("Reading Dynamic tracker...")
    dynamic_records = read_source(DYNAMIC_FILE)
    print(f"  Dynamic source rows: {len(dynamic_records)}")
    print()

    wb = load_workbook(TEMPLATE_FILE)

    required_sheets = ["App Inventory", "Scans", "Vulnerabilities"]
    missing_sheets = [s for s in required_sheets if s not in wb.sheetnames]

    if missing_sheets:
        raise ValueError(
            "new_template.xlsx is missing worksheet(s): "
            + ", ".join(missing_sheets)
        )

    # GLOBAL counters. They are NOT reset between Static and Dynamic.
    counters = {
        "app": 1,
        "scan": 1,
        "finding": 1,
    }

    static_result = migrate_one(
        static_records, wb, "static", counters
    )

    dynamic_result = migrate_one(
        dynamic_records, wb, "dynamic", counters
    )

    wb.save(OUTPUT_FILE)

    print()
    print("=" * 60)
    print("MIGRATION COMPLETED")
    print("=" * 60)
    print(f"Output: {OUTPUT_FILE}")
    print()
    print("Static:")
    print(f"  Apps  : {static_result['apps']}")
    print(f"  Scans : {static_result['scans']}")
    print()
    print("Dynamic:")
    print(f"  Apps  : {dynamic_result['apps']}")
    print(f"  Scans : {dynamic_result['scans']}")
    print()
    print("Next IDs:")
    print(f"  APP  : APP-{counters['app']:06d}")
    print(f"  SCAN : SCAN-{counters['scan']:06d}")
    print(f"  FIND : FIND-{counters['finding']:06d}")
    print()
    print("Done.")


if __name__ == "__main__":
    main()
